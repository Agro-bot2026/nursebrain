import os
import uuid
from app.config import UPLOAD_DIR

EXPORT_DIR = os.path.join(UPLOAD_DIR, "exports")
os.makedirs(EXPORT_DIR, exist_ok=True)


def generar_docx(titulo: str, contenido: str) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading(titulo, level=0)

    for parrafo in contenido.split("\n"):
        p = parrafo.rstrip()
        if not p:
            doc.add_paragraph("")
            continue
        doc.add_paragraph(p)

    doc.add_paragraph("")
    pie = doc.add_paragraph("Generado por NurseBrain AI - material de estudio.")
    for run in pie.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    nombre = f"{uuid.uuid4().hex}.docx"
    ruta = os.path.join(EXPORT_DIR, nombre)
    doc.save(ruta)
    return nombre


def generar_pdf(titulo: str, contenido: str) -> str:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    nombre = f"{uuid.uuid4().hex}.pdf"
    ruta = os.path.join(EXPORT_DIR, nombre)

    doc = SimpleDocTemplate(ruta, pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle("t", parent=estilos["Title"], fontSize=18,
                                   textColor=colors.HexColor("#6d28d9"))
    estilo_cuerpo = ParagraphStyle("c", parent=estilos["Normal"], fontSize=11, leading=16)
    estilo_pie = ParagraphStyle("p", parent=estilos["Normal"], fontSize=8, textColor=colors.grey)

    elementos = [Paragraph(_esc(titulo), estilo_titulo), Spacer(1, 0.5*cm)]
    for parrafo in contenido.split("\n"):
        p = parrafo.rstrip()
        if not p:
            elementos.append(Spacer(1, 0.25*cm))
        else:
            elementos.append(Paragraph(_esc(p), estilo_cuerpo))
    elementos.append(Spacer(1, 0.8*cm))
    elementos.append(Paragraph("Generado por NurseBrain AI - material de estudio.", estilo_pie))

    doc.build(elementos)
    return nombre


def _esc(t: str) -> str:
    return (t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
